# -*- coding: utf-8 -*-
r"""
mnemonic_tool.py - Tool nhập mẹo nhớ và ghi TRỰC TIẾP vào data/ai_mnemonics.json

Chạy:
    python tools/mnemonic_tool.py

Hoặc từ thư mục gốc:
    python tools/mnemonic_tool.py
"""

import os
import sys
import json
import re

# ═══════════════════════════════════════════════════════════════
#  ĐƯỜNG DẪN
# ═══════════════════════════════════════════════════════════════
# Cho phép chạy từ root hoặc từ tools/
if os.path.isfile("config.json"):
    ROOT = os.getcwd()
elif os.path.isfile(os.path.join("..", "config.json")):
    ROOT = os.path.abspath(os.path.join(os.getcwd(), ".."))
else:
    print("❌ Không tìm thấy config.json. Chạy tool từ thư mục gốc project.")
    sys.exit(1)

DATA_DIR = os.path.join(ROOT, "data")
os.makedirs(DATA_DIR, exist_ok=True)

MNEMONIC_FILE = os.path.join(DATA_DIR, "ai_mnemonics.json")
VOCAB_FILE = os.path.join(DATA_DIR, "tu_vung_hsk.xlsx")
FIXPY_FILE = os.path.join(DATA_DIR, "fixpy_datasets.json")


# ═══════════════════════════════════════════════════════════════
#  ĐỌC / GHI FILE JSON
# ═══════════════════════════════════════════════════════════════
def load_mnemonics():
    """Đọc file ai_mnemonics.json hiện có."""
    if not os.path.isfile(MNEMONIC_FILE):
        return {}
    try:
        with open(MNEMONIC_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
        if isinstance(data, dict):
            return data
        return {}
    except Exception as e:
        print("⚠️  Không đọc được file cũ: " + str(e))
        return {}


def save_mnemonics(data):
    """Ghi file ai_mnemonics.json (pretty print, giữ Unicode)."""
    try:
        with open(MNEMONIC_FILE, "w", encoding="utf-8") as f:
            json.dump(data, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print("❌ Lỗi ghi file: " + str(e))
        return False


# ═══════════════════════════════════════════════════════════════
#  ĐỌC DANH SÁCH TỪ VỰNG (từ fixpy_datasets.json hoặc Excel)
# ═══════════════════════════════════════════════════════════════
def load_vocab_list():
    """
    Trả về list [{'hsk': 'HSK1', 'stt': '4', 'zh': '爸爸', 'vi': 'Bố; ba', 'pinyin': 'bàba'}]
    """
    vocab = []

    # Ưu tiên đọc từ fixpy_datasets.json (nhanh hơn)
    if os.path.isfile(FIXPY_FILE):
        try:
            with open(FIXPY_FILE, "r", encoding="utf-8") as f:
                data = json.load(f)

            # Tìm dataset 'tu-vung'
            ds = data.get("tu-vung")
            if ds and isinstance(ds.get("data"), list):
                for row in ds["data"]:
                    vocab.append({
                        "hsk": str(row.get("hsk", "")).strip(),
                        "stt": str(row.get("stt_original", row.get("stt", ""))).strip(),
                        "zh": str(row.get("zh", "")).strip(),
                        "vi": str(row.get("vi", "")).strip(),
                        "pinyin": str(row.get("pinyin", "")).strip(),
                        "radical": row.get("radical"),
                        "topic": str(row.get("topic", "")).strip(),
                        "subject": str(row.get("subject", "")).strip(),
                    })
                print("✅ Đã load " + str(len(vocab)) + " từ từ fixpy_datasets.json")
                return vocab
        except Exception as e:
            print("⚠️  Lỗi đọc fixpy_datasets.json: " + str(e))

    # Fallback: đọc từ Excel
    if os.path.isfile(VOCAB_FILE):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(VOCAB_FILE, data_only=True, read_only=True)
            for sheet_name in wb.sheetnames:
                if not sheet_name.upper().startswith("HSK"):
                    continue
                hsk = sheet_name.strip()
                ws = wb[sheet_name]

                is_hsk79 = ("7" in hsk and "9" in hsk)
                COL_STT = 0
                COL_ZH = 1
                COL_PINYIN = 2
                COL_VI = 4 if is_hsk79 else 5

                for row in ws.iter_rows(min_row=3, values_only=True):
                    if not row or len(row) <= COL_ZH:
                        continue
                    zh = str(row[COL_ZH] or "").strip()
                    if not zh:
                        continue
                    vocab.append({
                        "hsk": hsk,
                        "stt": str(row[COL_STT] or "").strip(),
                        "zh": zh,
                        "vi": str(row[COL_VI] or "").strip() if COL_VI < len(row) else "",
                        "pinyin": str(row[COL_PINYIN] or "").strip() if COL_PINYIN < len(row) else "",
                        "radical": None,
                        "topic": "",
                        "subject": "",
                    })
            wb.close()
            print("✅ Đã load " + str(len(vocab)) + " từ từ Excel")
            return vocab
        except Exception as e:
            print("⚠️  Lỗi đọc Excel: " + str(e))

    return vocab


# ═══════════════════════════════════════════════════════════════
#  MENU CHÍNH
# ═══════════════════════════════════════════════════════════════
def print_banner():
    print("")
    print("=" * 62)
    print("  📚 MNEMONIC TOOL — Nhập mẹo nhớ HSK")
    print("  Ghi trực tiếp vào: " + MNEMONIC_FILE)
    print("=" * 62)
    print("")


def print_menu():
    print("")
    print("─" * 62)
    print("  CHỌN CHỨC NĂNG:")
    print("─" * 62)
    print("  1. ➕ Thêm mẹo nhớ mới (nhập tay đầy đủ)")
    print("  2. 📋 Xem danh sách từ CHƯA có mẹo nhớ")
    print("  3. ✏️  Sửa mẹo nhớ đã có")
    print("  4. 🗑️  Xóa mẹo nhớ")
    print("  5. 📊 Thống kê")
    print("  6. 🔍 Tìm kiếm")
    print("  0. 🚪 Thoát")
    print("─" * 62)


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 1: THÊM MỚI
# ═══════════════════════════════════════════════════════════════
def add_new_mnemonic(mnemonics, vocab):
    print("")
    print("─" * 62)
    print("  ➕ THÊM MẸO NHỚ MỚI")
    print("─" * 62)

    hsk = input("  HSK (VD: HSK1): ").strip().upper()
    if not hsk:
        print("  ❌ Hủy")
        return

    stt = input("  STT (VD: 4): ").strip()
    if not stt:
        print("  ❌ Hủy")
        return

    # Tìm chữ Hán trong vocab
    found = None
    for v in vocab:
        if v["hsk"].upper() == hsk and v["stt"] == stt:
            found = v
            break

    if found:
        print("")
        print("  ✅ Tìm thấy trong kho:")
        print("     Chữ Hán: " + found["zh"])
        print("     Nghĩa:   " + found["vi"])
        print("     Pinyin:  " + found["pinyin"])
        use_found = input("  Dùng thông tin này? (y/n): ").strip().lower()
        if use_found == "y":
            zh = found["zh"]
            vi = found["vi"]
            pinyin = found["pinyin"]
        else:
            zh = input("  Chữ Hán: ").strip()
            vi = input("  Nghĩa: ").strip()
            pinyin = input("  Pinyin: ").strip()
    else:
        print("  ⚠️  Không tìm thấy trong kho, nhập tay:")
        zh = input("  Chữ Hán (VD: 爸爸): ").strip()
        vi = input("  Nghĩa: ").strip()
        pinyin = input("  Pinyin: ").strip()

    if not zh:
        print("  ❌ Thiếu chữ Hán")
        return

    key = hsk + "|" + stt + "|" + zh

    # Kiểm tra tồn tại
    if key in mnemonics:
        print("  ⚠️  Key '" + key + "' đã tồn tại!")
        overwrite = input("  Ghi đè? (y/n): ").strip().lower()
        if overwrite != "y":
            return

    # Nhập bộ thủ
    print("")
    print("  ─── BỘ THỦ (tùy chọn, để trống bỏ qua) ───")
    print("  Format: 父|fù|4|Cha")
    radical = input("  Bộ thủ: ").strip()

    # Nhập mẹo nhớ (multi-line)
    print("")
    print("  ─── MẸO NHỚ (nhập nhiều dòng, gõ 'END' để kết thúc) ───")
    print("  VD:")
    print("    💡 Chi tiết từ: 父 (Cha) + 巴 (mong ước)")
    print("    🔗 Âm thanh: \"bàba\" nghe như \"ba ba\"")
    print("    📎 Ví dụ: 爸爸去哪儿了? (Bàba qù nǎr le?) - Bố đi đâu rồi?")
    print("")

    lines = []
    while True:
        line = input("  > ")
        if line.strip().upper() == "END":
            break
        lines.append(line)

    mnemonic = "\n".join(lines).strip()
    if not mnemonic:
        print("  ❌ Thiếu mẹo nhớ")
        return

    # Lưu
    mnemonics[key] = mnemonic
    if save_mnemonics(mnemonics):
        print("")
        print("  ✅ ĐÃ LƯU: " + key)
        print("  📁 File: " + MNEMONIC_FILE)
        print("  🔄 Chạy: python fix.py để build lại web")
    else:
        print("  ❌ Lỗi lưu file")


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 2: XEM DANH SÁCH CHƯA CÓ MẸO
# ═══════════════════════════════════════════════════════════════
def list_missing_mnemonics(mnemonics, vocab):
    print("")
    print("─" * 62)
    print("  📋 DANH SÁCH TỪ CHƯA CÓ MẸO NHỚ")
    print("─" * 62)

    if not vocab:
        print("  ❌ Không có dữ liệu từ vựng")
        return

    hsk_filter = input("  Lọc theo HSK (VD: HSK1, để trống = tất cả): ").strip().upper()

    missing = []
    for v in vocab:
        if hsk_filter and v["hsk"].upper() != hsk_filter:
            continue
        key = v["hsk"] + "|" + v["stt"] + "|" + v["zh"]
        if key not in mnemonics:
            missing.append(v)

    if not missing:
        print("  ✅ Không còn từ nào thiếu mẹo nhớ!")
        return

    print("")
    print("  Tìm thấy " + str(len(missing)) + " từ chưa có mẹo:")
    print("")
    print("  {:<8} {:<6} {:<15} {:<25}".format("HSK", "STT", "Chữ Hán", "Nghĩa"))
    print("  " + "─" * 60)

    for i, v in enumerate(missing[:50]):  # Chỉ hiện 50 đầu
        vi_short = v["vi"][:24] + "..." if len(v["vi"]) > 25 else v["vi"]
        print("  {:<8} {:<6} {:<15} {:<25}".format(
            v["hsk"], v["stt"], v["zh"], vi_short
        ))

    if len(missing) > 50:
        print("")
        print("  ... và " + str(len(missing) - 50) + " từ khác")

    print("")
    print("  💡 Tip: Dùng chức năng 1 để thêm mẹo nhớ cho các từ này")


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 3: SỬA
# ═══════════════════════════════════════════════════════════════
def edit_mnemonic(mnemonics):
    print("")
    print("─" * 62)
    print("  ✏️  SỬA MẸO NHỚ")
    print("─" * 62)

    if not mnemonics:
        print("  ❌ Chưa có mẹo nhớ nào")
        return

    # Hiện 30 key gần nhất
    keys = list(mnemonics.keys())[-30:]
    print("")
    print("  Danh sách " + str(len(keys)) + " key gần nhất:")
    for i, k in enumerate(keys, 1):
        print("  " + str(i) + ". " + k)

    print("")
    key_input = input("  Nhập key (hoặc số thứ tự): ").strip()

    # Nếu nhập số
    if key_input.isdigit():
        idx = int(key_input) - 1
        if 0 <= idx < len(keys):
            key = keys[idx]
        else:
            print("  ❌ Số không hợp lệ")
            return
    else:
        key = key_input

    if key not in mnemonics:
        print("  ❌ Không tìm thấy key: " + key)
        return

    print("")
    print("  ─── MẸO NHỚ HIỆN TẠI ───")
    print(mnemonics[key])
    print("  ─────────────────────────")
    print("")
    print("  Nhập mẹo nhớ MỚI (gõ 'END' để kết thúc, gõ 'CANCEL' để hủy):")

    lines = []
    while True:
        line = input("  > ")
        if line.strip().upper() == "END":
            break
        if line.strip().upper() == "CANCEL":
            print("  ❌ Hủy")
            return
        lines.append(line)

    new_mnemonic = "\n".join(lines).strip()
    if not new_mnemonic:
        print("  ❌ Rỗng, không lưu")
        return

    mnemonics[key] = new_mnemonic
    if save_mnemonics(mnemonics):
        print("  ✅ ĐÃ CẬP NHẬT: " + key)
        print("  🔄 Chạy: python fix.py")


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 4: XÓA
# ═══════════════════════════════════════════════════════════════
def delete_mnemonic(mnemonics):
    print("")
    print("─" * 62)
    print("  🗑️  XÓA MẸO NHỚ")
    print("─" * 62)

    if not mnemonics:
        print("  ❌ Chưa có mẹo nhớ nào")
        return

    key = input("  Nhập key cần xóa: ").strip()
    if key not in mnemonics:
        print("  ❌ Không tìm thấy key: " + key)
        return

    confirm = input("  Xác nhận xóa '" + key + "'? (y/n): ").strip().lower()
    if confirm != "y":
        print("  ❌ Hủy")
        return

    del mnemonics[key]
    if save_mnemonics(mnemonics):
        print("  ✅ ĐÃ XÓA: " + key)
        print("  🔄 Chạy: python fix.py")


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 5: THỐNG KÊ
# ═══════════════════════════════════════════════════════════════
def show_stats(mnemonics, vocab):
    print("")
    print("─" * 62)
    print("  📊 THỐNG KÊ")
    print("─" * 62)

    total_mnemonics = len(mnemonics)
    total_vocab = len(vocab)

    print("")
    print("  📁 File: " + MNEMONIC_FILE)
    if os.path.isfile(MNEMONIC_FILE):
        size_kb = os.path.getsize(MNEMONIC_FILE) / 1024
        print("  📦 Kích thước: " + "{:.1f}".format(size_kb) + " KB")
    print("")

    # Đếm theo HSK
    hsk_count = {}
    for key in mnemonics:
        parts = key.split("|")
        if len(parts) >= 1:
            hsk = parts[0]
            hsk_count[hsk] = hsk_count.get(hsk, 0) + 1

    # Đếm từ vựng theo HSK
    vocab_count = {}
    for v in vocab:
        hsk = v["hsk"]
        vocab_count[hsk] = vocab_count.get(hsk, 0) + 1

    print("  {:<10} {:<15} {:<15} {:<10}".format("HSK", "Có mẹo nhớ", "Tổng số từ", "Tỉ lệ"))
    print("  " + "─" * 54)

    all_hsk = sorted(set(list(hsk_count.keys()) + list(vocab_count.keys())))
    for hsk in all_hsk:
        have = hsk_count.get(hsk, 0)
        total = vocab_count.get(hsk, 0)
        pct = (have / total * 100) if total > 0 else 0
        print("  {:<10} {:<15} {:<15} {:.1f}%".format(hsk, have, total, pct))

    print("  " + "─" * 54)
    print("  {:<10} {:<15} {:<15}".format("TỔNG", total_mnemonics, total_vocab))


# ═══════════════════════════════════════════════════════════════
#  CHỨC NĂNG 6: TÌM KIẾM
# ═══════════════════════════════════════════════════════════════
def search_mnemonic(mnemonics):
    print("")
    print("─" * 62)
    print("  🔍 TÌM KIẾM")
    print("─" * 62)

    query = input("  Nhập từ khóa (key hoặc nội dung): ").strip()
    if not query:
        return

    query_lower = query.lower()
    results = []

    for key, value in mnemonics.items():
        if query_lower in key.lower() or query_lower in value.lower():
            results.append((key, value))

    if not results:
        print("  ❌ Không tìm thấy")
        return

    print("")
    print("  ✅ Tìm thấy " + str(len(results)) + " kết quả:")
    print("")

    for key, value in results[:10]:
        print("  ─── " + key + " ───")
        # Hiện 3 dòng đầu
        lines = value.split("\n")[:3]
        for line in lines:
            print("    " + line)
        print("")


# ═══════════════════════════════════════════════════════════════
#  MAIN LOOP
# ═══════════════════════════════════════════════════════════════
def main():
    print_banner()

    # Load dữ liệu
    print("  📂 Đang load...")
    mnemonics = load_mnemonics()
    vocab = load_vocab_list()

    print("  ✅ Mnemonics: " + str(len(mnemonics)) + " entries")
    print("  ✅ Vocab:     " + str(len(vocab)) + " từ")
    print("")

    while True:
        print_menu()
        choice = input("  👉 Chọn: ").strip()

        if choice == "1":
            add_new_mnemonic(mnemonics, vocab)
        elif choice == "2":
            list_missing_mnemonics(mnemonics, vocab)
        elif choice == "3":
            edit_mnemonic(mnemonics)
        elif choice == "4":
            delete_mnemonic(mnemonics)
        elif choice == "5":
            show_stats(mnemonics, vocab)
        elif choice == "6":
            search_mnemonic(mnemonics)
        elif choice == "0":
            print("")
            print("  👋 Tạm biệt!")
            break
        else:
            print("  ❌ Lựa chọn không hợp lệ")

        input("\n  [Enter] để tiếp tục...")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\n\n  👋 Đã thoát")
