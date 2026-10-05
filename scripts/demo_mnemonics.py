# -*- coding: utf-8 -*-
r"""
Test 3 từ với prompt nâng cấp - giữ format 5 dòng gọn gàng.
"""
import os
import sys
import json
import time
import unicodedata
from openai import OpenAI
import openpyxl


# ═══════════════════════════════════════════════════════════════════
#  CONFIG
# ═══════════════════════════════════════════════════════════════════
API_KEY = os.getenv("DASHSCOPE_API_KEY")

if not API_KEY:
    print("❌ Chưa set DASHSCOPE_API_KEY")
    sys.exit(1)

client = OpenAI(
    api_key=API_KEY,
    base_url="https://ws-8lcqaxwxdv0h9xzy.ap-southeast-1.maas.aliyuncs.com/compatible-mode/v1"
)

OUTPUT_FILE = "data/test_3_tu.json"
MODEL_ID = os.getenv("QWEN_MODEL", "qwen-plus-character")

DEMO_LIMIT = 3
DELAY_BETWEEN = 2.0


# ═══════════════════════════════════════════════════════════════════
#  TÌM FILE EXCEL
# ═══════════════════════════════════════════════════════════════════
def _norm(s):
    return unicodedata.normalize("NFC", s) if s else s


def find_vocab_excel():
    TARGET_NFC = "tu_vung_hsk.xlsx"
    folder = "data"
    if not os.path.isdir(folder):
        return None
    for fname in os.listdir(folder):
        if _norm(fname) == TARGET_NFC:
            return os.path.join(folder, fname)
    return None


EXCEL_FILE = find_vocab_excel()


# ═══════════════════════════════════════════════════════════════════
#  PROMPT NÂNG CẤP - GIỮ FORMAT 5 DÒNG
# ═══════════════════════════════════════════════════════════════════
def build_prompt(zh, pinyin, vi, hsk):
    return f"""Bạn là giáo viên tiếng Trung nổi tiếng, chuyên tạo mẹo nhớ khiến học trò Việt Nam "một lần nghe là nhớ mãi".

Từ cần nhớ: **{zh}** (pinyin: {pinyin}) - nghĩa: {vi} - HSK{hsk}

Hãy viết mẹo nhớ theo ĐÚNG định dạng 5 dòng sau (mỗi dòng 1 emoji):

💡 Chiết tự: [Phân tích bộ thủ và cách ghép. Nếu từ ghép, phân tích từng chữ rồi nêu quy tắc ghép nghĩa. PHẢI chỉ rõ: bộ nào + bộ nào = ý gì]
📌 Âm thanh: [Tìm âm Hán-Việt và so sánh với từ tiếng Việt gần âm. Nếu không có từ gần âm, tạo câu vè ngắn có vần]
🎬 Câu chuyện: [Một câu chuyện 2-3 câu có NHÂN VẬT cụ thể (chị Hoa, bác Ba, em Tí...), có HÀNH ĐỘNG, có CẢM XÚC. Câu chuyện phải liên kết logic các bộ thủ với nghĩa của từ. KHÔNG viết chung chung]
📎 Ví dụ: [Câu tiếng Trung ngắn + pinyin + nghĩa Việt]
🔗 Liên quan: [3-5 từ cùng bộ thủ, cùng âm, hoặc cùng chủ đề, kèm nghĩa]

QUY TẮC BẮT BUỘC:
- Giữ ĐÚNG 5 dòng, không thêm không bớt
- Mỗi dòng phải CỤ THỂ, có HÌNH ẢNH, tránh chung chung
- Câu chuyện phải có tên nhân vật và hành động rõ ràng
- Nếu có âm gần giống tiếng Việt: BẮT BUỘC dùng để liên tưởng
- Ưu tiên gắn với văn hóa Việt Nam (chợ quê, Tết, gia đình, trường học)

Ví dụ mẫu cho 菜 (cài - món ăn):
💡 Chiết tự: 艹 (bộ thảo - cỏ, rau) + 采 (bộ thủ - hái, ngắt) → Hái cỏ/rau về = MÓN ĂN (菜)
📌 Âm thanh: "cài" ≈ "cải" (rau cải). Câu vè: "Rau cải, rau cài, hái về nấu món ăn"
🎬 Câu chuyện: Bà Tư ở quê sáng sớm ra vườn, cúi xuống hái (采) những ngọn rau cỏ (艹) còn đọng sương. Bà mang về chế biến thành món ăn (菜) dâng lên bàn thờ tổ tiên ngày giỗ.
📎 Ví dụ: 这道菜很好吃 (Zhè dào cài hěn hǎo chī) - Món này rất ngon
🔗 Liên quan: 采 (cǎi - hái), 彩 (cǎi - màu sắc), 青菜 (qīngcài - rau xanh), 菜单 (càidān - thực đơn)

Bây giờ sinh mẹo cho: {zh}
Output (đúng 5 dòng):"""


# ═══════════════════════════════════════════════════════════════════
#  GỌI API
# ═══════════════════════════════════════════════════════════════════
def generate_mnemonic(zh, pinyin, vi, hsk):
    prompt = build_prompt(zh, pinyin, vi, hsk)
    for attempt in range(3):
        try:
            response = client.chat.completions.create(
                model=MODEL_ID,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.8,
                max_tokens=1000
            )
            if not response or not response.choices:
                time.sleep(3)
                continue
            msg = response.choices[0].message
            content = msg.content if hasattr(msg, 'content') else None
            if content and isinstance(content, str) and content.strip():
                return content.strip()
            time.sleep(3)
        except Exception as e:
            err = str(e)
            if "429" in err or "rate" in err.lower():
                wait = 15 * (attempt + 1)
                print(f"      ⏳ Rate limit, chờ {wait}s...")
                time.sleep(wait)
            else:
                print(f"      ⚠️  Lỗi API (lần {attempt+1}/3): {err}")
                time.sleep(3)
    return ""


# ═══════════════════════════════════════════════════════════════════
#  LOAD EXCEL
# ═══════════════════════════════════════════════════════════════════
def load_sample_words():
    if not EXCEL_FILE or not os.path.exists(EXCEL_FILE):
        print("❌ Không có file Excel vocab")
        return []
    wb = openpyxl.load_workbook(EXCEL_FILE, data_only=True, read_only=True)
    samples = []
    for sheet_name in wb.sheetnames:
        if len(samples) >= DEMO_LIMIT:
            break
        if not sheet_name.strip().upper().startswith("HSK"):
            continue
        ws = wb[sheet_name]
        is_hsk79 = "7-9" in sheet_name
        for row in ws.iter_rows(min_row=2, values_only=True):
            if len(samples) >= DEMO_LIMIT:
                break
            if not row or len(row) < 3:
                continue
            stt = str(row[0] or "").strip()
            zh = str(row[1] or "").strip()
            pinyin = str(row[2] or "").strip()
            vi_idx = 4 if is_hsk79 else 5
            vi = str(row[vi_idx] or "").strip() if len(row) > vi_idx else ""
            if not zh or not vi or len(zh) > 3:
                continue
            samples.append({
                "stt": stt, "zh": zh, "pinyin": pinyin, "vi": vi,
                "hsk": sheet_name.strip(),
                "key": f"{sheet_name.strip()}|{stt}|{zh}"
            })
    wb.close()
    return samples[:DEMO_LIMIT]


# ═══════════════════════════════════════════════════════════════════
#  MAIN
# ═══════════════════════════════════════════════════════════════════
def main():
    print("=" * 62)
    print("🧪 TEST 3 TỪ — PROMPT NÂNG CẤP (FORMAT 5 DÒNG)")
    print(f"   Model: {MODEL_ID}")
    print("=" * 62)
    if not EXCEL_FILE:
        print("❌ Không tìm thấy file Excel vocab. Dừng.")
        sys.exit(1)
    samples = load_sample_words()
    print(f"\n📚 Đã lấy {len(samples)} từ mẫu\n")
    if not samples:
        print("❌ Không có từ nào để test")
        sys.exit(1)
    results = {}
    for i, w in enumerate(samples, 1):
        print(f"\n{'='*62}")
        print(f"[{i}/{len(samples)}] {w['zh']} ({w['pinyin']}) - {w['vi']}")
        print(f"{'='*62}")
        mnemonic = generate_mnemonic(w["zh"], w["pinyin"], w["vi"], w["hsk"])
        if mnemonic:
            results[w["key"]] = mnemonic
            print(mnemonic)
        else:
            print("❌ FAIL")
        os.makedirs(os.path.dirname(OUTPUT_FILE), exist_ok=True)
        with open(OUTPUT_FILE, "w", encoding="utf-8") as f:
            json.dump(results, f, ensure_ascii=False, indent=2)
        time.sleep(DELAY_BETWEEN)
    print("\n" + "=" * 62)
    print(f"✅ HOÀN TẤT: {len(results)}/{len(samples)} từ")
    print(f"📁 File: {OUTPUT_FILE}")
    print("=" * 62)


if __name__ == "__main__":
    main()
